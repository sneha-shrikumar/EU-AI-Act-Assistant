"""Embedding-model and Annex III chunking diagnostic.

Two questions, both measured the same way as rank_diagnostic.py (how many of
the golden set's expected sources land in the top 8 by pure dense similarity):

  A. Does a bigger embedding model help? Re-embeds the CURRENT chunks (text
     exactly as stored in Chroma) with bge-small / bge-base / bge-large.
  B. Does splitting Annex III per area help? Replaces the 5 packed Annex III
     chunks with one chunk per numbered area (1-8), each under the Annex III
     heading, and re-ranks with bge-small.

Local and free. bge-large downloads ~1.3 GB on first run.

    python analysis/embedding_chunking_diagnostic.py [eval_all_<experiment>.xlsx]
"""
import re
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
from sentence_transformers import SentenceTransformer  # noqa: E402

import config  # noqa: E402
import eval_retrieval as er  # noqa: E402
from query import _get_collection  # noqa: E402

XLSX = sys.argv[1] if len(sys.argv) > 1 else BASE / "results" / "eval-sheets" / "eval_all_eu-ai-act-rag-e6e5c965.xlsx"

allc = _get_collection().get(include=["documents", "metadatas"])
ids, docs, metas = allc["ids"], allc["documents"], allc["metadatas"]

ws = openpyxl.load_workbook(XLSX)["all evals"]
h = [c.value for c in ws[1]]
questions = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(h, r))
    if r[0] and d["should_answer"] == "yes":
        exp = er.parse_refs(d["expected_source"] or "")
        if exp:
            questions.append((d["question"], exp))
n_refs = sum(len(e) for _, e in questions)


def score(model, corpus_docs, corpus_refs, label):
    E = model.encode(corpus_docs, normalize_embeddings=True, batch_size=16)
    hits, per_q = 0, []
    for q, exp in questions:
        qv = model.encode([config.QUERY_INSTRUCTION_PREFIX + q], normalize_embeddings=True)[0]
        order = np.argsort(-(E @ qv))
        ranks = {}
        for rank, i in enumerate(order, 1):
            for ref in corpus_refs[i] & exp:
                ranks.setdefault(ref, rank)
        hits += sum(1 for ref in exp if ranks.get(ref, 10**6) <= 8)
        per_q.append((q, {ref: ranks.get(ref) for ref in exp}))
    print(f"{label:<40} expected sources in top 8: {hits}/{n_refs}", flush=True)
    return per_q


# --- A. embedding models on the current chunks ------------------------------
refs = [er.parse_refs(m["citation"]) for m in metas]
results = {}
for name in ("BAAI/bge-small-en-v1.5", "BAAI/bge-base-en-v1.5", "BAAI/bge-large-en-v1.5"):
    results[name] = score(SentenceTransformer(name), docs, refs, name)

print("\nPer-reference ranks that changed (small -> base -> large):")
for (q, rs), (_, rb), (_, rl) in zip(*results.values()):
    for ref in sorted(rs):
        if (rs[ref], rb[ref], rl[ref]) != (rs[ref],) * 3:
            print(f"  {ref[0][:3]} {ref[1]:<4} {rs[ref]!s:>4} -> {rb[ref]!s:>4} -> {rl[ref]!s:>4} | {q[:60]}")

# --- B. Annex III split per area (bge-small) --------------------------------
annex3 = [i for i, m in enumerate(metas) if m["type"] == "Annex" and m["number"] == "III"]
body = "\n".join(
    # Strip each chunk's own heading lines and overlap carry-over so the areas
    # are rebuilt from body text only.
    "\n".join(line for line in docs[i].split("\n")
              if not line.startswith(("Annex III:", "High-risk AI systems pursuant", "… ")))
    for i in annex3
)
areas = re.split(r"(?m)^(?=\d\. )", body)
areas = [a.strip() for a in areas if re.match(r"\d\. ", a.strip())]
seen, split_docs = set(), []
for a in areas:
    num = a.split(".")[0]
    if num in seen:
        continue
    seen.add(num)
    split_docs.append("Annex III: High-risk AI systems referred to in Article 6(2)\n" + a)
print(f"\nAnnex III: {len(annex3)} packed chunks -> {len(split_docs)} per-area chunks")

keep = [i for i in range(len(docs)) if i not in set(annex3)]
docs_b = [docs[i] for i in keep] + split_docs
refs_b = [refs[i] for i in keep] + [{("annex", 3)}] * len(split_docs)
small = SentenceTransformer("BAAI/bge-small-en-v1.5")
res_b = score(small, docs_b, refs_b, "bge-small + Annex III split per area")

print("\nAnnex III rank, packed -> per-area (bge-small):")
for (q, ra), (_, rb) in zip(results["BAAI/bge-small-en-v1.5"], res_b):
    if ("annex", 3) in ra:
        print(f"  {ra[('annex', 3)]!s:>4} -> {rb[('annex', 3)]!s:>4} | {q[:65]}")
