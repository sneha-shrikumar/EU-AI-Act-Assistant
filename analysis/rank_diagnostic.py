"""Retrieval-options diagnostic: for every expected source in the golden set,
where does it rank under dense similarity, BM25, equal-weight hybrid (RRF),
and a bge-reranker-base cross-encoder over the top-50 dense / hybrid pool?

Free (local models only), but the reranker is slow on CPU (~35 s/question).
See docs/analysis/retrieval-options-analysis.md for the findings.

    python analysis/rank_diagnostic.py [eval_all_<experiment>.xlsx]
"""
import math
import re
import sys
from collections import Counter

from pathlib import Path
BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))
import numpy as np
import openpyxl

import eval_retrieval as er
from embeddings import embed_query
from query import _get_collection

col = _get_collection()
allc = col.get(include=["documents", "metadatas", "embeddings"])
ids, docs, metas = allc["ids"], allc["documents"], allc["metadatas"]
E = np.array(allc["embeddings"])
refs_of = [er.parse_refs(m["citation"]) for m in metas]

tok = lambda t: re.findall(r"[a-z0-9]+", t.lower())  # noqa: E731
D = [tok(d) for d in docs]
N = len(D)
avg = sum(map(len, D)) / N
df = Counter(w for d in D for w in set(d))
TF = [Counter(d) for d in D]


def bm25(q, k1=1.5, b=0.75):
    qs = tok(q)
    out = []
    for tf, d in zip(TF, D):
        s = 0.0
        for w in qs:
            if w in tf:
                idf = math.log(1 + (N - df[w] + 0.5) / (df[w] + 0.5))
                s += idf * tf[w] * (k1 + 1) / (tf[w] + k1 * (1 - b + b * len(d) / avg))
        out.append(s)
    return np.array(out)


def rank_of(scores, ref):
    order = np.argsort(-scores)
    for r, i in enumerate(order, 1):
        if ref in refs_of[i]:
            return r
    return None


def rrf(*score_lists, k=60):
    tot = np.zeros(N)
    for s in score_lists:
        order = np.argsort(-s)
        for r, i in enumerate(order, 1):
            tot[i] += 1 / (k + r)
    return tot


from sentence_transformers import CrossEncoder  # noqa: E402
ce = CrossEncoder("BAAI/bge-reranker-base")

ws = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else BASE / "results" / "eval-sheets" / "eval_all_eu-ai-act-rag-342c66bc.xlsx")["all evals"]
h = [c.value for c in ws[1]]
rows = [dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

print(f"{'ref':12} dense  bm25  hybrid  rerank(dense50)  rerank(hybrid50) | question")
agg = {k: [] for k in ("dense", "bm25", "hybrid", "rr_d", "rr_h")}
for d in rows:
    if d["should_answer"] != "yes":
        continue
    exp = er.parse_refs(d["expected_source"] or "")
    if not exp:
        continue
    q = d["question"]
    dense = E @ np.array(embed_query(q))
    bm = bm25(q)
    hy = rrf(dense, bm)
    ranks = {}
    for name, s in (("dense", dense), ("bm25", bm), ("hybrid", hy)):
        ranks[name] = {ref: rank_of(s, ref) for ref in exp}
    for name, base in (("rr_d", dense), ("rr_h", hy)):
        pool = list(np.argsort(-base)[:50])
        sc = ce.predict([(q, docs[i]) for i in pool])
        s = np.full(N, -1e9)
        for i, v in zip(pool, sc):
            s[i] = v
        ranks[name] = {ref: rank_of(s, ref) for ref in exp}
    for ref in sorted(exp):
        vals = [ranks[n][ref] for n in agg]
        for n, v in zip(agg, vals):
            agg[n].append(v)
        fmt = lambda v: f"{v:>5}" if v is not None and v < 1e8 else "  -  "  # noqa: E731
        print(f"{ref[0][:3]} {ref[1]:<8} " + "  ".join(fmt(v) for v in vals) + f" | {q[:55]}")

print()
for n, v in agg.items():
    hit8 = sum(1 for x in v if x is not None and x <= 8)
    print(f"{n:7} refs in top 8: {hit8}/{len(v)}")
