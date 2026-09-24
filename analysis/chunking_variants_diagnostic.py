"""Chunking-strategy diagnostic: scores four whole-corpus chunking variants.

See plans/structural-chunking-plan.md. Each variant sets config.CHUNK_PACKING
and config.OVERLAP_SCOPE, re-chunks the PDF (extracted once), embeds the chunks
in memory with bge-small (no Chroma write), and ranks every golden question by
pure dense similarity -- the same method as embedding_chunking_diagnostic.py.

    A  size packing,     document overlap   (today)
    B  size packing,     window overlap
    C  no packing,       window overlap     (one chunk per structural unit)
    D  semantic packing, window overlap

Local and free.

    python analysis/chunking_variants_diagnostic.py [eval_all_<experiment>.xlsx]
"""
import statistics
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

import chunking  # noqa: E402
import config  # noqa: E402
import eval_retrieval as er  # noqa: E402
from embeddings import embed_passages, embed_query  # noqa: E402
from ingest import extract_pdf_text  # noqa: E402

XLSX = sys.argv[1] if len(sys.argv) > 1 else BASE / "results" / "eval-sheets" / "eval_all_eu-ai-act-rag-e6e5c965.xlsx"
K = config.TOP_K

VARIANTS = [
    ("A. Today", "size", "document"),
    ("B. Today + overlap fix", "size", "window"),
    ("C. Structural", "none", "window"),
    ("D. Semantic", "semantic", "window"),
]

# Informational only -- these never decide the winner (see the plan).
TARGETS = [
    ("deployer", "who is a deployer", ("article", 3)),
    ("dated", "what is this document dated", ("preamble", 0)),
    ("#38 Annex III", "when is the AI system considered high risk", ("annex", 3)),
    ("CV screening", "CV screening", ("annex", 3)),
    ("Italy", "Children in Italy", ("annex", 3)),
]

ws = openpyxl.load_workbook(XLSX)["all evals"]
h = [c.value for c in ws[1]]
questions = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(h, r))
    if r[0] and d["should_answer"] == "yes":
        exp = er.parse_refs(d["expected_source"] or "")
        if exp:
            questions.append((d["question"], exp))
n_refs = sum(len(e) for _, e in questions)
qvecs = np.asarray([embed_query(q) for q, _ in questions])
print(f"{len(questions)} questions, {n_refs} expected sources\n", flush=True)

raw = extract_pdf_text(config.PDF_PATH)


def run(label, packing, overlap):
    config.CHUNK_PACKING, config.OVERLAP_SCOPE = packing, overlap
    chunks = chunking.chunk_document(raw)
    docs = [c["text"] for c in chunks]
    refs = [er.parse_refs(c["citation"]) for c in chunks]
    E = np.asarray(embed_passages(docs))
    tokens = [chunking.count_tokens(d) for d in docs]

    hits = {k: 0 for k in (1, 3, 5, K)}
    distinct, ranks_by_q = [], []
    for qv, (_q, exp) in zip(qvecs, questions):
        order = np.argsort(-(E @ qv))
        ranks = {}
        for rank, i in enumerate(order, 1):
            for ref in refs[i] & exp:
                ranks.setdefault(ref, rank)
        for k in hits:
            hits[k] += sum(1 for ref in exp if ranks.get(ref, 10**6) <= k)
        distinct.append(len(set().union(*(refs[i] for i in order[:K]))))
        ranks_by_q.append(ranks)

    return {
        "label": label,
        "top8": hits[K],
        "r1": hits[1], "r3": hits[3], "r5": hits[5],
        "distinct": statistics.mean(distinct),
        "chunks": len(chunks),
        "median_tok": statistics.median(tokens),
        "max_tok": max(tokens),
        "ranks": ranks_by_q,
    }


results = []
for label, packing, overlap in VARIANTS:
    res = run(label, packing, overlap)
    results.append(res)
    print(f"{label:<24} top-{K}: {res['top8']}/{n_refs}  chunks: {res['chunks']}", flush=True)

print(f"\n{'Variant':<24} {'top-8':>6} {'R@1':>4} {'R@3':>4} {'R@5':>4} "
      f"{'distinct':>8} {'chunks':>6} {'med tok':>7} {'max tok':>7}")
for r in results:
    print(f"{r['label']:<24} {r['top8']:>6} {r['r1']:>4} {r['r3']:>4} {r['r5']:>4} "
          f"{r['distinct']:>8.2f} {r['chunks']:>6} {r['median_tok']:>7} {r['max_tok']:>7}")

print("\nTarget ranks (informational):")
print(f"{'target':<16}" + "".join(f"{r['label'][:2]:>6}" for r in results))
for name, needle, ref in TARGETS:
    qi = next(i for i, (q, _) in enumerate(questions) if needle.lower() in q.lower())
    print(f"{name:<16}" + "".join(f"{r['ranks'][qi].get(ref, '-')!s:>6}" for r in results))

# --- Decision rule, fixed before running (see the plan) ----------------------
a = results[0]
print()
if a["top8"] != 37:
    print(f"SANITY CHECK FAILED: variant A scored {a['top8']}, expected 37. Stop.")
    sys.exit(1)
best = max(results, key=lambda r: (r["top8"], -r["chunks"]))
if best["top8"] >= a["top8"] + 3:
    winner = best
elif results[1]["top8"] >= a["top8"]:
    winner = results[1]
else:
    winner = a
print(f"Winner: {winner['label']} ({winner['top8']}/{n_refs})")
assert all(r["max_tok"] <= config.CHUNK_TOKENS for r in results), "chunk over budget"
