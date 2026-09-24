"""Agentic-retrieval ladder: how much does each agent stage add to retrieval?

See plans/agentic-rag-plan.md. For every answerable golden question, runs
agent.gather_evidence (stages 1-4, NO answer generation, NO judges) at each
rung of the ladder and scores the evidence exactly like the chunking
diagnostics: expected sources among the first 8 citations, out of 61.

    rung 0  baseline   query._retrieve, top 8 (no LLM)
    rung 1  analyse    rewrite + decompose + RRF fusion
    rung 2  reflect    + coverage check and follow-up searches
    rung 3  select     + LLM evidence selection

Each LLM rung runs RUNS times (LLM output varies slightly even at
temperature 0) and the mean decides. Costs a few hundred Haiku calls.

    python analysis/agentic_retrieval_diagnostic.py [eval_all_<experiment>.xlsx]
"""
import os
import statistics
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

os.environ["LANGSMITH_TRACING"] = "false"  # diagnostic only; keep LangSmith clean
BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(BASE / ".env")
os.environ["LANGSMITH_TRACING"] = "false"

import openpyxl  # noqa: E402

import agent  # noqa: E402
import config  # noqa: E402
import eval_retrieval as er  # noqa: E402
import query  # noqa: E402

XLSX = sys.argv[1] if len(sys.argv) > 1 else BASE / "results" / "eval-sheets" / "eval_all_eu-ai-act-rag-275ca613.xlsx"
K = config.TOP_K
RUNS = 2
WORKERS = 4
MARGIN = 3  # fixed before running, same as the chunking decision

ws = openpyxl.load_workbook(XLSX)["all evals"]
h = [c.value for c in ws[1]]
questions = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(h, r))
    if r[0] and d["should_answer"] == "yes":
        exp = er.parse_refs(d["expected_source"] or "")
        if exp:
            questions.append((d["question"], exp))
n_refs = sum(len(e) for _, e in questions)
print(f"{len(questions)} questions, {n_refs} expected sources\n", flush=True)


def score(citation_lists):
    """citation_lists[i] = ranked citations for questions[i]."""
    hits = {k: 0 for k in (1, 3, 5, K)}
    distinct, per_q = [], []
    for cits, (_q, exp) in zip(citation_lists, questions):
        refs = [er.parse_refs(c) for c in cits]
        ranks = {}
        for rank, rs in enumerate(refs, 1):
            for ref in rs & exp:
                ranks.setdefault(ref, rank)
        for k in hits:
            hits[k] += sum(1 for ref in exp if ranks.get(ref, 10**6) <= k)
        distinct.append(len(set().union(*refs[:K])) if refs else 0)
        per_q.append(sum(1 for ref in exp if ranks.get(ref, 10**6) <= K))
    return {"top8": hits[K], "r1": hits[1], "r3": hits[3], "r5": hits[5],
            "distinct": statistics.mean(distinct), "per_q": per_q}


def run_rung(stages):
    def one(q):
        g = agent.gather_evidence(q, K, stages=stages)
        return g
    with ThreadPoolExecutor(WORKERS) as pool:
        gathered = list(pool.map(one, [q for q, _ in questions]))
    res = score([[c["citation"] for c in g["evidence"]] for g in gathered])
    res["calls"] = statistics.mean(g["llm_calls"] for g in gathered)
    res["fallbacks"] = sum(bool(g["fallbacks"]) for g in gathered)
    res["dropped_refs"] = sum(len(g["trace"]["dropped_refs"]) for g in gathered)
    return res


# --- rung 0: today's retrieval ------------------------------------------------
query._get_collection()
baseline = score([[c["citation"] for c in query._retrieve(q, K)] for q, _ in questions])
baseline.update(calls=0, fallbacks=0, dropped_refs=0)
rows = [("0 baseline", [baseline])]
print(f"{'0 baseline':<12} top-8: {baseline['top8']}/{n_refs}", flush=True)

for stages in ("analyse", "reflect", "select"):
    runs = []
    for i in range(RUNS):
        res = run_rung(stages)
        runs.append(res)
        print(f"{stages:<12} run {i + 1}: top-8 {res['top8']}/{n_refs}  "
              f"calls/q {res['calls']:.1f}  fallbacks {res['fallbacks']}", flush=True)
    rows.append((f"{('analyse', 'reflect', 'select').index(stages) + 1} {stages}", runs))


def mean(runs, key):
    return statistics.mean(r[key] for r in runs)


print(f"\n{'Rung':<12} {'top-8':>12} {'mean':>6} {'R@1':>5} {'R@3':>5} {'R@5':>5} "
      f"{'distinct':>8} {'calls/q':>7} {'fallbk':>6} {'dropped refs':>12}")
for label, runs in rows:
    each = "/".join(str(r["top8"]) for r in runs)
    print(f"{label:<12} {each:>12} {mean(runs, 'top8'):>6.1f} {mean(runs, 'r1'):>5.1f} "
          f"{mean(runs, 'r3'):>5.1f} {mean(runs, 'r5'):>5.1f} {mean(runs, 'distinct'):>8.2f} "
          f"{mean(runs, 'calls'):>7.1f} {mean(runs, 'fallbacks'):>6.1f} "
          f"{mean(runs, 'dropped_refs'):>12.1f}")

print("\nPer question, expected sources in top 8 (mean over runs):")
print(f"{'':<62}" + "".join(f"{label.split()[1][:7]:>8}" for label, _ in rows))
for i, (q, exp) in enumerate(questions):
    vals = [statistics.mean(r["per_q"][i] for r in runs) for _, runs in rows]
    if len(set(vals)) > 1:
        print(f"{q[:58]:<58} /{len(exp):<2} " + "".join(f"{v:>8.1f}" for v in vals))

# --- decision rule, fixed before running -------------------------------------
kept_label, kept = rows[0][0], mean(rows[0][1], "top8")
for label, runs in rows[1:]:
    m = mean(runs, "top8")
    if m >= kept + MARGIN:
        kept_label, kept = label, m
print(f"\nKept rung: {kept_label} ({kept:.1f}/{n_refs})")
