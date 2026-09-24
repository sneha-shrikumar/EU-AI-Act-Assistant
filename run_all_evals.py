"""One command for every eval: run the golden dataset, score it, write one sheet.

    python run_all_evals.py                         # full run (41 questions)
    python run_all_evals.py --limit 3               # cheap smoke test
    python run_all_evals.py --system agentic        # agent.py's staged pipeline
    python run_all_evals.py --experiment NAME       # re-score an existing run,
                                                    # no new answers, no cost
    python run_all_evals.py --import-labels FILE.xlsx
        # you edited judge scores in an eval_all_*.xlsx: record your
        # corrections in evals/human_labels.json (the judges learn from
        # them on the next run) and recalculate that sheet's scores. No cost.

Steps:
  1. Offline self-tests (provision_refs + eval_retrieval parsers). Free; stops
     here if any fail.
  2. evals/run_langsmith_eval.py: every question through the RAG pipeline,
     scored by the three LLM judges (groundedness, answer relevance, answer
     correctness). This is the only step that costs money.
  3. Retrieval eval (eval_retrieval.py logic): strict pass/fail, recall@k and
     precision@k at k = 1, 3, 5, 8, plus the top-3 citation-ranking hit.
  4. Writes eval_all_<experiment>.xlsx: one row per question with every
     score, plus a summary sheet, and prints the summary.
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from dotenv import load_dotenv

import eval_retrieval as er
from eval_citation_ranking import score_ranking_row
from evals.judges import HUMAN_LABELS_PATH, load_human_labels

BASE_DIR = Path(__file__).resolve().parent
JUDGE_KEYS = ("groundedness", "answer_relevance", "answer_correctness")


def _run(cmd: list[str]) -> str:
    """Run a python step, streaming its output live; returns captured stdout."""
    print(f"\n$ {' '.join(cmd)}", flush=True)
    proc = subprocess.Popen(
        [sys.executable, *cmd], cwd=BASE_DIR, stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace",
    )
    out = []
    for line in proc.stdout:
        print(line, end="", flush=True)
        out.append(line)
    if proc.wait() != 0:
        raise SystemExit(f"Step failed: {' '.join(cmd)}")
    return "".join(out)


def _score_experiment(experiment_name: str) -> tuple[list[dict], dict]:
    from langsmith import Client

    client = Client()
    experiment, rows = er.fetch_experiment(client, er.DATASET_NAME, experiment_name)
    paragraph_index = er.build_paragraph_index()

    judge, why = {}, {}
    for fb in client.list_feedback(run_ids=[run.id for run, _ in rows]):
        judge.setdefault(fb.run_id, {})[fb.key] = fb.score
        why.setdefault(fb.run_id, {})[fb.key] = fb.comment or ""

    results = []
    for run, reference in rows:
        question = (run.inputs or {}).get("question", "")
        expected_source = reference.get("expected_source") or ""
        should_answer = (reference.get("should_answer") or "yes").strip().lower()
        outputs = run.outputs or {}
        citations = outputs.get("citations") or []
        refused = bool(outputs.get("refused"))

        expected_refs = er.parse_refs(expected_source)
        citation_refs = set().union(*(er.parse_refs(c) for c in citations)) if citations else set()
        score, note = er.score_row(
            expected_refs, citation_refs, citations, should_answer, refused,
            expected_source, question, paragraph_index,
        )
        rank, _ = score_ranking_row(expected_refs, citations, should_answer, expected_source)
        latency = ((run.end_time - run.start_time).total_seconds()
                   if run.end_time and run.start_time else None)
        results.append({
            "question": question, "expected_source": expected_source,
            "should_answer": should_answer, "citations": citations,
            "refused": refused, "retrieval": score, "note": note or (run.error or "")[:200],
            "top3": rank,
            "pr": er.precision_recall_at_k(expected_refs, citations)
                  if should_answer == "yes" else {},
            "judges": dict(judge.get(run.id, {})),
            "judge_raw": dict(judge.get(run.id, {})),  # never overridden
            "judge_why": why.get(run.id, {}),
            "overridden": set(),
            "latency": latency, "answer": outputs.get("answer") or "",
            "expected_answer": reference.get("expected_answer") or "",
            "errored": bool(run.error),
        })
    return results, {"experiment": experiment.name}


def _apply_human_labels(results: list[dict], experiment: str) -> int:
    """Replace judge scores with the human's wherever evals/human_labels.json
    has a correction for this experiment, so rebuilding a sheet never
    silently reverts a reviewed score. Returns how many were applied."""
    labels = {(lab["question"], lab["key"]): lab["human_score"]
              for lab in load_human_labels() if lab.get("experiment") == experiment}
    applied = 0
    for r in results:
        for key in JUDGE_KEYS:
            if (r["question"], key) in labels:
                r["judges"][key] = labels[(r["question"], key)]
                r["overridden"].add(key)
                applied += 1
    return applied


def _import_labels(xlsx: Path) -> str:
    """Diff the judge columns of an edited eval_all_*.xlsx against the judge's
    own LangSmith scores and record every difference in human_labels.json.
    Returns the experiment name. Re-importing the same sheet is idempotent:
    an existing (experiment, question, key) entry is updated, not duplicated,
    and its `lesson` is kept."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    experiment = next((row[1] for row in wb["summary"].iter_rows(values_only=True)
                       if row and row[0] == "experiment"), None)
    if not experiment:
        raise SystemExit(f"No 'experiment' row in the summary sheet of {xlsx}.")
    ws = wb["all evals"]
    header = [c.value for c in ws[1]]
    edited = {row[0]: dict(zip(header, row))
              for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}

    results, _ = _score_experiment(experiment)
    labels = load_human_labels()
    index = {(lab.get("experiment"), lab["question"], lab["key"]): lab for lab in labels}
    changes = 0
    for r in results:
        row = edited.get(r["question"])
        if row is None:
            continue
        for key in JUDGE_KEYS:
            cell = row.get(key)
            human = None if cell in (None, "") else int(cell)
            judge = r["judge_raw"].get(key)
            judge = None if judge is None else int(judge)
            ident = (experiment, r["question"], key)
            if human == judge:
                # Edited back to agree with the judge: drop any stale label.
                if ident in index:
                    labels.remove(index.pop(ident))
                    changes += 1
                continue
            if human is None:
                continue  # blanking a cell is not a verdict
            entry = index.get(ident)
            if entry is None:
                entry = {"experiment": experiment, "question": r["question"],
                         "key": key, "lesson": ""}
                labels.append(entry)
                index[ident] = entry
            entry.update({"judge_score": judge, "human_score": human,
                          "judge_reasoning": r["judge_why"].get(key, "")})
            changes += 1
            print(f"  {key}: judge {judge} -> human {human} | {r['question'][:70]}")

    HUMAN_LABELS_PATH.write_text(json.dumps(labels, indent=2, ensure_ascii=False),
                                 encoding="utf-8")
    backup = xlsx.with_name(f"{xlsx.stem}_before-import{xlsx.suffix}")
    shutil.copy2(xlsx, backup)
    print(f"{changes} label change(s) -> {HUMAN_LABELS_PATH}"
          f"  (your sheet backed up to {backup.name})")
    return experiment


def _mean(values):
    values = [v for v in values if v is not None]
    return (sum(values) / len(values), len(values)) if values else (None, 0)


def _summary(results: list[dict]) -> list[tuple[str, float | None, int]]:
    rows = []
    for key in JUDGE_KEYS:
        rows.append((key, *_mean([r["judges"].get(key) for r in results])))
    if any(r["overridden"] for r in results):
        for key in JUDGE_KEYS:
            rows.append((f"{key} (judge only, before human review)",
                         *_mean([r["judge_raw"].get(key) for r in results])))
        reviewed = sum(1 for r in results for k in JUDGE_KEYS
                       if r["judges"].get(k) is not None)
        overruled = sum(len(r["overridden"]) for r in results)
        rows.append(("judge-human agreement", (reviewed - overruled) / reviewed, reviewed))
    rows.append(("retrieval (strict)", *_mean([r["retrieval"] for r in results])))
    rows.append(("top-3 citation hit", *_mean([r["top3"] for r in results])))
    for k in er.K_VALUES:
        rows.append((f"recall@{k}", *_mean([r["pr"].get(f"recall@{k}") for r in results])))
    for k in er.K_VALUES:
        rows.append((f"precision@{k} (lower bound)",
                     *_mean([r["pr"].get(f"precision@{k}") for r in results])))
    rows.append(("latency s (mean)", *_mean([r["latency"] for r in results])))
    return rows


def _write_excel(path: Path, experiment: str, results: list[dict], summary) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "all evals"
    headers = (["question", "expected_source", "should_answer", "citations",
                "refused", "groundedness", "answer_relevance", "answer_correctness",
                "retrieval (strict)", "top-3 hit"]
               + [f"R@{k}" for k in er.K_VALUES] + [f"P@{k}" for k in er.K_VALUES]
               + ["retrieval note", "latency s", "expected_answer", "answer"])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    zero = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    human = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    blank = lambda v: "" if v is None else v  # noqa: E731
    for r in results:
        pr = r["pr"]
        ws.append(
            [r["question"], r["expected_source"], r["should_answer"],
             ", ".join(r["citations"]), "yes" if r["refused"] else "no"]
            + [blank(r["judges"].get(k)) for k in JUDGE_KEYS]
            + [blank(r["retrieval"]), blank(r["top3"])]
            + [round(pr[f"recall@{k}"], 2) if pr else "" for k in er.K_VALUES]
            + [round(pr[f"precision@{k}"], 2) if pr else "" for k in er.K_VALUES]
            + [r["note"], round(r["latency"], 1) if r["latency"] else "",
                  r["expected_answer"], r["answer"]]
        )
        for cell in ws[ws.max_row]:
            cell.alignment = wrap
            if cell.column in range(6, 11) and cell.value == 0:
                cell.fill = zero
            if cell.column in range(6, 9) and JUDGE_KEYS[cell.column - 6] in r["overridden"]:
                cell.fill = human
    for col, width in {"A": 45, "B": 22, "D": 35, "S": 30, "U": 60, "V": 70}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"

    ss = wb.create_sheet("summary")
    ss.append(["metric", "score", "n"])
    for cell in ss[1]:
        cell.font = Font(bold=True)
    for name, value, n in summary:
        ss.append([name, round(value, 3) if value is not None else "", n])
    ss.append([])
    ss.append(["experiment", experiment])
    ss.append(["LEGEND", "blue judge cell = human override (evals/human_labels.json); "
                         "pink = score 0"])
    ss.append(["NOTE", "Precision is a lower bound: the golden expected_source lists "
                       "only primary sources. All averages are per-question (macro)."])
    ss.column_dimensions["A"].width = 30
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--experiment", default=None,
                    help="Re-score an existing experiment instead of running a new one.")
    ap.add_argument("--limit", type=int, default=None, help="Only the first N questions.")
    ap.add_argument("--concurrency", type=int, default=4)
    ap.add_argument("--system", choices=("rag", "agentic", "full_doc"), default="rag",
                    help="Which pipeline answers the questions (see "
                         "evals/run_langsmith_eval.py).")
    ap.add_argument("--import-labels", type=Path, default=None, metavar="XLSX",
                    help="Record your edited judge scores from an eval_all_*.xlsx "
                         "and recalculate it.")
    args = ap.parse_args()
    load_dotenv(BASE_DIR / ".env")

    if args.import_labels:
        args.experiment = _import_labels(args.import_labels)

    _run(["provision_refs.py"])
    _run(["eval_retrieval.py", "--selftest"])

    experiment = args.experiment
    if experiment is None:
        cmd = ["evals/run_langsmith_eval.py", "--concurrency", str(args.concurrency),
               "--system", args.system]
        if args.limit is not None:
            cmd += ["--limit", str(args.limit)]
        match = re.search(r"Done\. Experiment: (\S+)", _run(cmd))
        if not match:
            raise SystemExit("Could not find the experiment name in the eval output.")
        experiment = match.group(1)

    print(f"\nScoring retrieval + collecting judge scores for {experiment} ...")
    results, meta = _score_experiment(experiment)
    applied = _apply_human_labels(results, meta["experiment"])
    if applied:
        print(f"Applied {applied} human-reviewed score(s) from {HUMAN_LABELS_PATH.name}.")
    summary = _summary(results)
    out = BASE_DIR / "results" / "eval-sheets" / f"eval_all_{meta['experiment']}.xlsx"
    _write_excel(out, meta["experiment"], results, summary)

    errored = sum(r["errored"] for r in results)
    print(f"\n=== {meta['experiment']}  ({len(results)} questions"
          f"{f', {errored} ERRORED' if errored else ''}) ===")
    for name, value, n in summary:
        shown = f"{value:.3f}" if value is not None else "n/a"
        print(f"  {name:<28} {shown:>7}   (n={n})")
    print(f"\nWritten to: {out}")


if __name__ == "__main__":
    main()
