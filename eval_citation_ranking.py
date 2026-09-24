"""Citation-ranking judge: was the right chunk retrieved near the top?

Fetches a finished experiment from LangSmith (same source as
eval_retrieval.py -- the runs produced by query.answer_question over the
"EU AI Acts evaluation" dataset) and checks whether at least one expected
article/annex/recital reference appears among the first 3 citations the
pipeline returned. This is a ranking signal on top of eval_retrieval.py's
plain recall check: a chunk that was retrieved but buried past position 3
still scores 0 here.

Scoring, at article/annex/recital granularity (paragraphs ignored, so a
citation "Article 13(2)" satisfies an expected "article 13"):
  - should_answer=yes: 1 if at least one expected reference appears among the
    top-3 citations, else 0.
  - should_answer=no (out of scope): left blank (N/A) -- refusal correctness
    is eval_retrieval.py's job, not this judge's.
  - should_answer=yes but no expected_source recorded: left blank (N/A).

Usage:
    python eval_citation_ranking.py                     # latest experiment
    python eval_citation_ranking.py --experiment NAME   # a specific experiment
"""
import argparse

from dotenv import load_dotenv

from eval_retrieval import DATASET_NAME, fetch_experiment, parse_refs


def score_ranking_row(expected_refs, citations, should_answer, expected_source):
    """Returns (score, note) where score is 1, 0, or None for N/A."""
    if should_answer != "yes":
        return None, "N/A -- ranking judge only scores should_answer=yes"
    if not expected_refs:
        if expected_source.strip():
            return None, "N/A -- no article/annex/recital in expected source"
        return None, "N/A -- no expected source recorded"

    top3_refs = set()
    for citation in citations[:3]:
        top3_refs |= parse_refs(citation)

    if expected_refs & top3_refs:
        return 1, ""
    return 0, "no expected ref in top 3"


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

def write_excel(path: str, experiment_name: str, results: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "citation ranking"

    headers = [
        "question (input)", "expected_source", "top-3 citations (output)",
        "should_answer", "score", "note",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    miss_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in results:
        ws.append([
            r["question"], r["expected_source"], ", ".join(r["top3_citations"]),
            r["should_answer"],
            r["score"] if r["score"] is not None else "",
            r["note"],
        ])
        row = ws[ws.max_row]
        for cell in row:
            cell.alignment = wrap
        if r["score"] == 0:
            for cell in row:
                cell.fill = miss_fill

    scored = [r for r in results if r["score"] is not None]
    hits = sum(1 for r in scored if r["score"] == 1)
    ws.append([])
    ws.append(["experiment", experiment_name])
    ws.append(["examples", len(results)])
    ws.append(["scored", len(scored)])
    ws.append(["top-3 hits", hits])
    ws.append(["hit rate", f"{hits / len(scored):.0%}" if scored else "n/a"])
    for row in ws.iter_rows(min_row=ws.max_row - 4):
        row[0].font = Font(bold=True)

    widths = {"A": 45, "B": 25, "C": 35, "D": 12, "E": 8, "F": 30}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def evaluate(dataset_name: str, experiment_name: str | None, out_path: str | None):
    load_dotenv()
    from langsmith import Client

    client = Client()
    experiment, rows = fetch_experiment(client, dataset_name, experiment_name)

    results = []
    for run, reference in rows:
        question = (run.inputs or {}).get("question", "")
        expected_source = reference.get("expected_source") or ""
        should_answer = (reference.get("should_answer") or "yes").strip().lower()
        citations = (run.outputs or {}).get("citations") or []

        expected_refs = parse_refs(expected_source)
        score, note = score_ranking_row(
            expected_refs, citations, should_answer, expected_source
        )
        results.append({
            "question": question, "expected_source": expected_source,
            "top3_citations": citations[:3], "all_citations": citations,
            "should_answer": should_answer, "score": score, "note": note,
        })

    out_path = out_path or f"eval_ranking_{experiment.name}.xlsx"
    write_excel(out_path, experiment.name, results)

    scored = [r for r in results if r["score"] is not None]
    hits = sum(1 for r in scored if r["score"] == 1)
    print(f"Experiment:   {experiment.name}")
    print(f"Examples:     {len(results)}  (scored: {len(scored)})")
    if scored:
        print(f"Top-3 hit rate: {hits}/{len(scored)} = {hits / len(scored):.0%} "
              f"(an expected reference was among the first 3 citations)")
    print(f"Written to: {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dataset", default=DATASET_NAME)
    parser.add_argument("--experiment", default=None,
                        help="experiment name; defaults to the latest one")
    parser.add_argument("--out", default=None, help="output .xlsx path")
    args = parser.parse_args()

    evaluate(args.dataset, args.experiment, args.out)


if __name__ == "__main__":
    main()
