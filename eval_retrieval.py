"""Retrieval-recall evaluator: did the tool retrieve the right chunks?

Fetches a finished experiment from LangSmith (the runs produced by
query.answer_question over the "EU AI Acts evaluation" dataset), compares each
run's retrieved-chunk citations against the golden expected_source, and writes
an Excel sheet with a 1/0 score per question.

Scoring, at article/annex/recital granularity (paragraphs ignored, so a
citation "Article 13(2)" satisfies an expected "article 13"):
  - should_answer=yes: 1 only if ALL expected references appear among the
    run's citations (missing even one scores 0).
  - should_answer=yes AND expected_source is a single bare article/annex (no
    paragraph named, no other refs mixed in) AND the question itself names
    that same article/annex (e.g. "summarize article 112", "what does
    article 97 talk about"): additionally 0 unless EVERY paragraph-chunk that
    provision has in the live corpus was retrieved (checked against the
    Chroma collection query.py itself reads). Narrow questions that merely
    resolve to one article without naming it in the question text (e.g.
    "what is an AI system" -> article 3) are NOT held to full-article
    coverage -- only the "summarize this whole provision" pattern is.
  - should_answer=no (out of scope): 1 if the tool refused, else 0.
  - should_answer=yes but no expected_source recorded: left blank (N/A).

Alongside that strict pass/fail, the sheet reports FRACTIONAL precision@k and
recall@k at k = 1, 3, 5, 8 (see precision_recall_at_k). The strict score answers
"was coverage complete?"; these answer "how complete, and at what cost in
precision?" -- which the strict score cannot show, and which matters because
recall rises monotonically with TOP_K while precision does not. Raising TOP_K
therefore always looks like a win under the strict score alone.

CAVEAT ON PRECISION: the golden expected_source lists only the PRIMARY sources
for a question. A retrieved chunk that genuinely helps but is not listed there
is counted as irrelevant, so reported precision is a LOWER BOUND, not true
precision. Recall is unaffected. This bites hardest on broad scenario questions
whose answer legitimately draws on provisions the golden row never enumerates.

Usage:
    python eval_retrieval.py                     # latest experiment on the dataset
    python eval_retrieval.py --experiment NAME   # a specific experiment
    python eval_retrieval.py --selftest          # run parser checks only
"""
import argparse
import re
import sys

from dotenv import load_dotenv

DATASET_NAME = "EU AI Acts evaluation"

# ---------------------------------------------------------------------------
# Reference parsing
# ---------------------------------------------------------------------------

# Keyword stems are deliberately loose: the golden expected_source is
# hand-typed and contains variants/typos like "artcile 97". "Section 4" is
# skipped via the None kind -- sections are not a chunk citation type, and the
# golden rows that mention one always list the concrete article too.
# "Preamble" is a singleton -- chunking.py emits exactly one such chunk
# (citation "Preamble", no number attached) -- so unlike the other kinds it
# doesn't wait for a following number token; see parse_refs.
KIND_STEMS = [
    (re.compile(r"^art", re.IGNORECASE), "article"),
    (re.compile(r"^recit", re.IGNORECASE), "recital"),
    (re.compile(r"^annex", re.IGNORECASE), "annex"),
    (re.compile(r"^pream", re.IGNORECASE), "preamble"),
    (re.compile(r"^section", re.IGNORECASE), None),
]

ROMAN_RE = re.compile(r"^[IVXLC]+$", re.IGNORECASE)
ROMAN_VALUES = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100}

# A token is either a kind keyword or a number reference. Number references may
# be arabic ("13"), roman annex numerals ("III"), or an arabic range ("9-15").
TOKEN_RE = re.compile(
    r"[A-Za-z]+|\d+\s*[-–]\s*\d+|\d+",
)

# Paragraph/point suffixes attached to a provision number -- "10(2)", "13(a)",
# "Annex III(2)". Stripped before tokenizing so the suffix number is never
# mistaken for another provision (matching is article-level by design).
PARA_SUFFIX_RE = re.compile(r"(?<=[0-9IVXLCivxlc])\s*\([^)]*\)")


def _roman_to_int(s: str) -> int:
    total = 0
    prev = 0
    for ch in reversed(s.upper()):
        val = ROMAN_VALUES[ch]
        total = total - val if val < prev else total + val
        prev = max(prev, val)
    return total


def parse_refs(text: str) -> set[tuple[str, int]]:
    """Parses free-form legal references into {(kind, number), ...}.

    Handles both the strict chunk-citation format ("Article 6(2)",
    "Annex III(1)") and the loose golden format ("Recital 24, 25, article 2",
    "Article 9-15, 16", "annex 3", "artcile 97"). Bare numbers inherit the kind
    of the last keyword seen; paragraph suffixes like "(2)"/"(a)" are stripped
    up front (see PARA_SUFFIX_RE).
    """
    refs: set[tuple[str, int]] = set()
    if not text:
        return refs
    text = PARA_SUFFIX_RE.sub("", text)

    kind: str | None = None
    kind_active = False  # False after "Section": swallow its numbers silently
    for token in TOKEN_RE.findall(text):
        if token[0].isalpha():
            matched = False
            for stem_re, stem_kind in KIND_STEMS:
                if stem_re.match(token):
                    kind, kind_active = stem_kind, True
                    matched = True
                    if stem_kind == "preamble":
                        # Singleton -- there's only ever one, always "(preamble,
                        # 0)" regardless of what number (if any) follows in the
                        # loose golden text or whether one follows at all (the
                        # real chunk citation is bare "Preamble", no number).
                        refs.add(("preamble", 0))
                    break
            if not matched and ROMAN_RE.match(token) and kind == "annex" and kind_active:
                refs.add(("annex", _roman_to_int(token)))
            # Other words ("what", "and", single "(a)" letters) are ignored;
            # they do not reset the active kind, so "Recital 24 and 25" works.
            continue

        if not kind_active or kind is None or kind == "preamble":
            continue
        if "-" in token or "–" in token:
            start, end = re.split(r"\s*[-–]\s*", token)
            refs.update((kind, n) for n in range(int(start), int(end) + 1))
        else:
            refs.add((kind, int(token)))
    return refs


# Strict chunk-citation format only (never the loose golden format): used to
# recover the paragraph a citation points at, which parse_refs deliberately
# discards (it matches at article level only). "Article 112(8)" -> paragraph
# "8"; "Article 4" (no internal split) -> paragraph None.
# Plural forms and ranges are both required: chunking.py packs short
# consecutive provisions into one chunk, citing them as "Recitals 1-2" or
# "Article 70(1-5)". The singular-only pattern matched neither, so packed
# chunks silently dropped out of the completeness check entirely.
CITATION_RE = re.compile(
    r"^(Article|Annex|Recital)s?\s+([A-Za-z0-9]+(?:\s*[-–]\s*\d+)?)"
    r"(?:\(([^)]+)\))?$"
)


def expand_range(label: str) -> set[str]:
    """Every individual label a packed metadata/citation span covers.

    "1-5" -> {"1","2","3","4","5"}; "2" -> {"2"}; "a" -> {"a"}. Both dash
    characters are accepted so a stale en-dash corpus still scores correctly."""
    if not label:
        return set()
    parts = re.split(r"\s*[-–]\s*", label.strip())
    if len(parts) == 2 and all(p.isdigit() for p in parts):
        return {str(n) for n in range(int(parts[0]), int(parts[1]) + 1)}
    return {label.strip()}


def _citation_detail(citation: str) -> tuple[str, int, set[str]] | None:
    """Returns (kind, number, paragraphs) for a strict chunk citation, or None
    if it doesn't match that format.

    `paragraphs` is a SET, because one packed citation can cover several:
    "Article 70(1-5)" contributes five paragraph labels, not the single opaque
    label "1-5" (which would never match the index built per-paragraph)."""
    m = CITATION_RE.match(citation.strip())
    if not m:
        return None
    kind_word, number_raw, paragraph = m.groups()
    kind = kind_word.lower()
    number_raw = re.split(r"\s*[-–]\s*", number_raw)[0]
    if kind == "annex" and ROMAN_RE.match(number_raw):
        number = _roman_to_int(number_raw)
    else:
        try:
            number = int(number_raw)
        except ValueError:
            return None
    return kind, number, expand_range(paragraph)


def build_paragraph_index() -> dict[tuple[str, int], set[str]]:
    """Returns {(kind, number): {paragraph_label, ...}} for every
    Article/Annex in the live corpus, read straight from the Chroma
    collection query.py itself reads (so it always matches what's actually
    retrievable right now). Provisions kept whole -- no internal paragraph
    split -- are omitted; there's nothing to check completeness against."""
    import chromadb

    import config

    client = chromadb.PersistentClient(path=config.CHROMA_DIR)
    try:
        collection = client.get_collection(config.COLLECTION_NAME)
    except Exception as exc:
        raise SystemExit(
            f"Collection '{config.COLLECTION_NAME}' not found at "
            f"{config.CHROMA_DIR}. Run `python ingest.py` first."
        ) from exc
    metadatas = collection.get(include=["metadatas"])["metadatas"]

    index: dict[tuple[str, int], set[str]] = {}
    for meta in metadatas:
        kind = (meta.get("type") or "").lower()
        paragraph = meta.get("paragraph")
        if kind not in ("article", "annex") or not paragraph:
            continue
        number_raw = (meta.get("number") or "").split("-")[0]
        if kind == "annex" and ROMAN_RE.match(number_raw):
            number = _roman_to_int(number_raw)
        else:
            try:
                number = int(number_raw)
            except ValueError:
                continue  # e.g. "47a" -- not handled anywhere else either
        # A packed chunk's paragraph metadata is a span ("1-5"); the index must
        # hold the individual paragraphs so completeness compares like with
        # like against the expanded citation labels.
        index.setdefault((kind, number), set()).update(expand_range(paragraph))
    return index


def score_row(expected_refs, citation_refs, citations, should_answer, refused,
              expected_source, question, paragraph_index):
    """Returns (score, note) where score is 1, 0, or None for N/A."""
    if should_answer == "no":
        return (1 if refused else 0), "out of scope: scored on refusal"
    if not expected_refs:
        if expected_source.strip():
            return None, "N/A -- no article/annex/recital/preamble in expected source"
        return None, "N/A -- no expected source recorded"
    if not (expected_refs <= citation_refs):
        missing = expected_refs - citation_refs
        return 0, f"missing: {', '.join(f'{k} {n}' for k, n in sorted(missing))}"

    # Whole-provision completeness: only for a single bare article/annex ref
    # (no paragraph already named in the golden source, no other refs mixed
    # in) whose number the QUESTION ITSELF names ("summarize article 112",
    # "what does article 97 talk about") -- i.e. the question is about one
    # entire provision, not a narrow topic that merely happens to be answered
    # somewhere inside a (possibly huge) article, like "what is an AI system"
    # -> article 3, which should not be held to full-article coverage.
    if (len(expected_refs) == 1 and not PARA_SUFFIX_RE.search(expected_source)
            and expected_refs <= parse_refs(question)):
        kind, number = next(iter(expected_refs))
        corpus_paragraphs = paragraph_index.get((kind, number))
        if corpus_paragraphs and len(corpus_paragraphs) > 1:
            covered = set()
            for c in citations:
                detail = _citation_detail(c)
                if detail and (detail[0], detail[1]) == (kind, number):
                    covered |= detail[2]
            missing_paragraphs = corpus_paragraphs - covered
            if missing_paragraphs:
                missing_sorted = sorted(
                    missing_paragraphs,
                    key=lambda p: (0, int(p)) if p.isdigit() else (1, p),
                )
                return 0, (
                    f"incomplete retrieval: missing {kind} {number} "
                    f"paragraph(s) {', '.join(missing_sorted)} "
                    f"({len(covered)}/{len(corpus_paragraphs)} retrieved)"
                )
    return 1, ""


# ---------------------------------------------------------------------------
# Fractional precision@k / recall@k
# ---------------------------------------------------------------------------

K_VALUES = (1, 3, 5, 8)


def precision_recall_at_k(expected_refs, citations, k_values=K_VALUES) -> dict:
    """Fractional precision@k and recall@k for one question.

    Returns {"recall@1": float, ..., "precision@5": float}, or an empty dict
    when there is nothing to score (no expected refs).

    Unlike score_row's strict all-or-nothing check, these are fractions: a
    question expecting {Annex III, Article 9} that retrieved only Article 9
    scores recall 0.5 here and 0 there. Both are reported, because the strict
    score answers "was coverage complete?" and these answer "how complete, and
    at what cost in precision?".

    RANK ORDER IS LOAD-BEARING. `citations` arrives ranked by similarity --
    query.py builds it from _retrieve's return, which is Chroma's ranked
    result -- so citations[0] is rank 1 and the [:k] slices mean what they say.
    Anything that reorders citations upstream silently invalidates every @k
    number here.

    Two asymmetries that are deliberate, not bugs:

    * Recall counts PROVISIONS, precision counts CHUNKS. One packed citation
      ("Recitals 1-2") expands to two refs and can satisfy two expected refs
      at once for recall, but it is still a single retrieved item for
      precision.
    * Retrieving three chunks of one expected article counts as three relevant
      items (P@5 = 0.6), not one. That is textbook precision@k -- each
      retrieved item is judged independently -- and keeps these numbers
      comparable to published benchmarks.
    """
    if not expected_refs:
        return {}

    # Per-citation ref sets, computed once; index i == rank i+1.
    per_citation = [parse_refs(c) for c in citations]

    out = {}
    for k in k_values:
        top = per_citation[:k]
        found = set().union(*top) & expected_refs if top else set()
        out[f"recall@{k}"] = len(found) / len(expected_refs)
        # min(k, len(citations)) rather than bare k, so a run that returned
        # fewer than k chunks is not penalised for slots it never filled.
        denom = min(k, len(citations))
        relevant = sum(1 for refs in top if refs & expected_refs)
        out[f"precision@{k}"] = (relevant / denom) if denom else 0.0
    return out


def macro_average(results: list[dict], key: str) -> float | None:
    """Mean of `key` across rows that have it -- a MACRO average (every question
    weighted equally), not micro (pooling hit counts across questions). The two
    differ whenever questions have different numbers of expected refs; macro is
    reported because the unit of interest here is the question, not the ref."""
    vals = [r["pr"][key] for r in results if r.get("pr")]
    return sum(vals) / len(vals) if vals else None


# ---------------------------------------------------------------------------
# LangSmith fetch
# ---------------------------------------------------------------------------

def fetch_experiment(client, dataset_name: str, experiment_name: str | None):
    """Returns (experiment_project, rows) sourced entirely from the LangSmith
    experiment itself: each row pairs a run (input + outputs) with its
    reference output (fetched by the run's reference_example_id), which is
    where the golden expected_source lives. The dataset is only used to look
    up which project is the (latest) experiment -- no golden CSV involved."""
    dataset = client.read_dataset(dataset_name=dataset_name)
    experiments = list(client.list_projects(reference_dataset_id=dataset.id))
    if not experiments:
        raise SystemExit(f"No experiments found for dataset '{dataset_name}'.")

    if experiment_name:
        matches = [e for e in experiments if e.name == experiment_name]
        if not matches:
            names = ", ".join(e.name for e in experiments)
            raise SystemExit(
                f"Experiment '{experiment_name}' not found. Available: {names}"
            )
        experiment = matches[0]
    else:
        experiment = max(experiments, key=lambda e: e.start_time)

    runs = [
        r for r in client.list_runs(project_id=experiment.id, is_root=True)
        if r.reference_example_id
    ]
    runs.sort(key=lambda r: r.start_time)

    example_ids = [r.reference_example_id for r in runs]
    reference_by_example_id = {
        ex.id: (ex.outputs or {})
        for ex in client.list_examples(example_ids=example_ids)
    } if example_ids else {}

    rows = [(run, reference_by_example_id.get(run.reference_example_id, {})) for run in runs]
    return experiment, rows


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

def write_excel(path: str, experiment_name: str, results: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "retrieval recall"

    headers = [
        "question (input)", "expected_source", "citations (output)",
        "answer", "refused", "score", "note",
    ] + [f"R@{k}" for k in K_VALUES] + [f"P@{k}" for k in K_VALUES]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    miss_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in results:
        pr = r.get("pr") or {}
        ws.append([
            r["question"], r["expected_source"], ", ".join(r["citations"]),
            r["answer"], "yes" if r["refused"] else "no",
            r["score"] if r["score"] is not None else "",
            r["note"],
        ] + [round(pr[f"recall@{k}"], 2) if pr else "" for k in K_VALUES]
          + [round(pr[f"precision@{k}"], 2) if pr else "" for k in K_VALUES])
        row = ws[ws.max_row]
        for cell in row:
            cell.alignment = wrap
        if r["score"] == 0:
            for cell in row:
                cell.fill = miss_fill

    scored = [r for r in results if r["score"] is not None]
    misses = sum(1 for r in scored if r["score"] == 0)
    ws.append([])
    ws.append(["experiment", experiment_name])
    ws.append(["examples", len(results)])
    ws.append(["scored", len(scored)])
    ws.append(["failures (score 0)", misses])
    ws.append(["failure rate", f"{misses / len(scored):.0%}" if scored else "n/a"])
    summary_start = ws.max_row - 4

    rows_with_pr = [r for r in results if r.get("pr")]
    if rows_with_pr:
        ws.append([])
        ws.append(["macro-averaged precision / recall",
                   f"over {len(rows_with_pr)} question(s)"])
        for k in K_VALUES:
            ws.append([
                f"recall@{k} (macro)", round(macro_average(results, f"recall@{k}"), 3),
                f"precision@{k} (macro)", round(macro_average(results, f"precision@{k}"), 3),
            ])
        ws.append([])
        ws.append(["NOTE", "Precision is a LOWER BOUND: the golden expected_source "
                           "lists only the primary sources, so a genuinely useful "
                           "retrieved chunk that is not listed counts as irrelevant. "
                           "Recall is unaffected."])
        ws.append(["NOTE", "Averages are MACRO (each question weighted equally), "
                           "not micro (pooled hit counts)."])

    for row in ws.iter_rows(min_row=summary_start):
        row[0].font = Font(bold=True)

    widths = {"A": 45, "B": 25, "C": 35, "D": 60, "E": 8, "F": 8, "G": 30,
              "H": 7, "I": 7, "J": 7, "K": 7, "L": 7, "M": 7, "N": 7, "O": 7}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def evaluate(dataset_name: str, experiment_name: str | None, out_path: str | None):
    load_dotenv()
    from langsmith import Client

    client = Client()
    experiment, rows = fetch_experiment(client, dataset_name, experiment_name)
    paragraph_index = build_paragraph_index()

    results = []
    for run, reference in rows:
        question = (run.inputs or {}).get("question", "")
        expected_source = reference.get("expected_source") or ""
        should_answer = (reference.get("should_answer") or "yes").strip().lower()

        outputs = run.outputs or {}
        citations = outputs.get("citations") or []
        refused = bool(outputs.get("refused"))

        expected_refs = parse_refs(expected_source)
        citation_refs = set()
        for citation in citations:
            citation_refs |= parse_refs(citation)

        score, note = score_row(
            expected_refs, citation_refs, citations, should_answer, refused,
            expected_source, question, paragraph_index,
        )
        # Scored on the same rows as `score`: an out-of-scope question has no
        # relevant set, so precision/recall are undefined rather than zero.
        pr = (precision_recall_at_k(expected_refs, citations)
              if should_answer == "yes" else {})
        results.append({
            "question": question, "expected_source": expected_source,
            "citations": citations, "answer": outputs.get("answer") or "",
            "refused": refused, "score": score, "note": note, "pr": pr,
        })

    out_path = out_path or f"eval_results_{experiment.name}.xlsx"
    write_excel(out_path, experiment.name, results)

    scored = [r for r in results if r["score"] is not None]
    misses = sum(1 for r in scored if r["score"] == 0)
    print(f"Experiment:   {experiment.name}")
    print(f"Examples:     {len(results)}  (scored: {len(scored)})")
    if scored:
        print(f"Failure rate: {misses}/{len(scored)} = {misses / len(scored):.0%} "
              f"(retrieval did not cover the expected source)")
    rows_with_pr = [r for r in results if r.get("pr")]
    if rows_with_pr:
        print(f"Macro-averaged over {len(rows_with_pr)} question(s):")
        for k in K_VALUES:
            r_at = macro_average(results, f"recall@{k}")
            p_at = macro_average(results, f"precision@{k}")
            print(f"  k={k}:  recall@{k} = {r_at:.2f}   precision@{k} = {p_at:.2f}")
        print("  (precision is a LOWER BOUND -- see module docstring)")
    print(f"Written to: {out_path}")


def selftest() -> None:
    cases = [
        ("article 13", {("article", 13)}),
        ("Article 13(a)", {("article", 13)}),
        ("article 10(2), 10(5)", {("article", 10)}),
        ("Recital 24, 25, article 2",
         {("recital", 24), ("recital", 25), ("article", 2)}),
        ("annex 3, Article 9-15, 16",
         {("annex", 3)} | {("article", n) for n in range(9, 16)} | {("article", 16)}),
        ("Section 4, article 28", {("article", 28)}),
        ("artcile 97", {("article", 97)}),
        ("Annex III(2)", {("annex", 3)}),
        ("Annex VIII", {("annex", 8)}),
        ("Article 71, Annex 8", {("article", 71), ("annex", 8)}),
        ("preamble 0", {("preamble", 0)}),
        ("Preamble", {("preamble", 0)}),
        ("", set()),
    ]
    failed = False
    for text, expected in cases:
        got = parse_refs(text)
        status = "ok " if got == expected else "FAIL"
        if got != expected:
            failed = True
        print(f"[{status}] {text!r} -> {sorted(got)}")
    # --- precision@k / recall@k ---------------------------------------
    # Synthetic citation lists, ranked. Each case pins one property of the
    # definitions so a future refactor can't quietly change the semantics.
    pr_cases = [
        # (label, expected_source, citations, {metric: value})
        ("no hits at all", "article 9",
         ["Recital 1", "Recital 2", "Recital 3", "Recital 4", "Recital 5"],
         {"recall@5": 0.0, "precision@5": 0.0}),
        ("single hit at rank 1", "article 9",
         ["Article 9(1)", "Recital 2", "Recital 3", "Recital 4", "Recital 5"],
         {"recall@1": 1.0, "precision@1": 1.0, "precision@5": 0.2}),
        ("hit at rank 4 only -- a RANKING failure, not a retrieval one",
         "article 9",
         ["Recital 1", "Recital 2", "Recital 3", "Article 9(1)", "Recital 5"],
         {"recall@1": 0.0, "recall@3": 0.0, "recall@5": 1.0, "precision@5": 0.2}),
        ("partial coverage scores a FRACTION (strict score would be 0)",
         "annex 3, article 9",
         ["Article 9(1)", "Recital 2", "Recital 3", "Recital 4", "Recital 5"],
         {"recall@5": 0.5, "precision@5": 0.2}),
        ("one PACKED chunk satisfies two expected refs: recall 1.0 from a "
         "single retrieved item, which precision still counts once",
         "recital 1, recital 2",
         ["Recitals 1-2", "Recital 9", "Recital 8", "Recital 7", "Recital 6"],
         {"recall@5": 1.0, "precision@5": 0.2}),
        ("three chunks of one expected article count as THREE relevant items",
         "article 9",
         ["Article 9(1)", "Article 9(2)", "Article 9(3)", "Recital 4", "Recital 5"],
         {"recall@5": 1.0, "precision@5": 0.6, "precision@3": 1.0}),
        ("fewer citations than k -- denominator is min(k, len), not k",
         "article 9",
         ["Article 9(1)", "Recital 2"],
         {"precision@5": 0.5, "recall@5": 1.0}),
    ]
    for label, expected_source, citations, wants in pr_cases:
        got = precision_recall_at_k(parse_refs(expected_source), citations)
        bad = {m: (got.get(m), v) for m, v in wants.items()
               if abs(got.get(m, -1) - v) > 1e-9}
        # Recall must never decrease as k grows.
        mono = all(got[f"recall@{a}"] <= got[f"recall@{b}"] + 1e-9
                   for a, b in zip(K_VALUES, K_VALUES[1:]))
        if bad or not mono:
            failed = True
            print(f"[FAIL] {label}: mismatches={bad} monotonic_recall={mono}")
        else:
            print(f"[ok ] {label}")

    if failed:
        sys.exit(1)
    print("All parser checks passed.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dataset", default=DATASET_NAME)
    parser.add_argument("--experiment", default=None,
                        help="experiment name; defaults to the latest one")
    parser.add_argument("--out", default=None, help="output .xlsx path")
    parser.add_argument("--selftest", action="store_true",
                        help="run parser checks and exit")
    args = parser.parse_args()

    if args.selftest:
        selftest()
        return
    evaluate(args.dataset, args.experiment, args.out)


if __name__ == "__main__":
    main()
